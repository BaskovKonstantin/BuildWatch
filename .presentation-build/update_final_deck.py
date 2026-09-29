"""Refresh docs/BuildWatch-ЛЦТ2026-final.pptx in place for the ЛЦТ submission."""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "docs" / "BuildWatch-ЛЦТ2026-final.pptx"
SHOTS = ROOT / "context" / "ui_shots"
DIAGRAMS = ROOT / "docs" / "diagrams"

DARK = RGBColor(0x1F, 0x1A, 0x2A)
PURPLE = RGBColor(0x4E, 0x1A, 0x7A)
ACCENT = RGBColor(0xFF, 0x00, 0x53)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"


def shape(slide, name, text_start=None):
    for sh in slide.shapes:
        if sh.name != name:
            continue
        if text_start is None or (sh.has_text_frame and sh.text_frame.text.startswith(text_start)):
            return sh
    raise KeyError(f"{name!r} / {text_start!r}")


def remove(sh):
    sh._element.getparent().remove(sh._element)


def set_bullet(paragraph, char="•", color=ACCENT):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.18)))
    pPr.set("indent", str(-Inches(0.18)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr")).set("val", str(color))
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def add_run(paragraph, text, size, color=DARK, bold=False):
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def fill(tf, items, size=12, color=DARK, lead_color=PURPLE, bullets=True, space=6, heading=None,
         heading_size=16, heading_color=PURPLE, subheading=None):
    """items: list of str or (bold_lead, rest)."""
    tf.clear()
    tf.word_wrap = True
    first = tf.paragraphs[0]
    paragraphs = []
    if heading:
        add_run(first, heading, heading_size, heading_color, bold=True)
        first.space_after = Pt(2 if subheading else 8)
        first = None
        if subheading:
            sub = tf.add_paragraph()
            add_run(sub, subheading, 11, ACCENT, bold=True)
            sub.space_after = Pt(10)
    for item in items:
        p = first if first is not None else tf.add_paragraph()
        first = None
        if isinstance(item, tuple):
            add_run(p, item[0], size, lead_color, bold=True)
            add_run(p, item[1], size, color)
        else:
            add_run(p, item, size, color)
        p.space_after = Pt(space)
        if bullets:
            set_bullet(p)
        paragraphs.append(p)
    return paragraphs


def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0.05))
    return tb


def replace_picture(slide, old, image_path):
    left, top, width, height = old.left, old.top, old.width, old.height
    remove(old)
    pic = slide.shapes.add_picture(str(image_path), left, top, width=width)
    if pic.height > height:
        remove(pic)
        pic = slide.shapes.add_picture(str(image_path), left, top, height=height)
        pic.left = left + (width - pic.width) // 2
    else:
        pic.top = top + (height - pic.height) // 2
    return pic


def picture_fit(slide, path, x, y, max_w, max_h):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(max_h))
    if pic.width > Inches(max_w):
        remove(pic)
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(max_w))
    line = pic.line
    line.color.rgb = RGBColor(0xDC, 0xD2, 0xEF)
    line.width = Pt(0.75)
    return pic


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def set_text_keep_style(sh, text):
    runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def team_card(slide, name_role, lines):
    box = shape(slide, "Текст 8", name_role)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, (text, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        add_run(p, text, 11 if i else 12.5, color, bold=bold)
        p.space_after = Pt(1)


def slide_team(s):
    team_card(s, "Разработчик", [
        ("Разработчик, капитан", True, DARK),
        ("Telegram: @KonstantBas", False, DARK),
        ("Тел.: +7 918 318-47-80", False, DARK),
    ])
    team_card(s, "Продуктовый менеджер", [
        ("Продуктовый менеджер", True, DARK),
        ("Telegram: @ddmanakov", False, DARK),
        ("Тел.: +7 960 000-15-17", False, DARK),
    ])


def slide_solution(s):
    set_text_keep_style(shape(s, "Текст 14"), "Техническая суть решения")
    right_title = shape(s, "Текст 16")
    set_text_keep_style(right_title, "Маркетинговая суть решения")
    right_run = right_title.text_frame.paragraphs[0].runs[0]
    right_run.font.name = None
    right_run.font.size = Pt(20)
    fill(shape(s, "Текст 2").text_frame, [
        ("Снимок → этап плана. ", "Кадр с датой привязывается к объекту и этапу календарного графика."),
        ("YOLOv8m + SAHI. ", "Детектор находит 8 классов техники, правила сверяют её с этапом."),
        ("Сигнал с объяснением. ", "Инспектор подтверждает или отклоняет вывод."),
    ], size=11, space=4)
    pipeline = {f"Прямоугольник: скругленные углы {n}" for n in (27, 29, 31, 35, 38)}
    pipeline |= {f"Прямоугольник {n}" for n in (28, 30, 32, 36, 39)}
    pipeline |= {"Стрелка: вправо 33", "Стрелка: вправо 34",
                 "Прямая соединительная линия 37", "Прямая соединительная линия 40"}
    for sh in list(s.shapes):
        if sh.name in pipeline:
            remove(sh)
    tb = textbox(s, 7.1, 2.1, 5.4, 4.3)
    fill(tb.text_frame, [
        ("Для кого: ", "инспекторы стройконтроля ДГП, технадзор заказчика и девелоперы."),
        ("Ценность: ", "вместо ручного просмотра камер — очередь сигналов с доказательством на кадре."),
        ("Эффект: ", "раннее выявление отставания по графику и простоя техники, меньше выездов."),
        ("Внедрение: ", "работает на существующих камерах площадок, без нового оборудования."),
        ("Развитие: ", "база знаний объекта, сметы, переписка, распознавание в реальном времени."),
    ], size=12, space=7)
    set_notes(s, "Слайд повторяет строгий шаблон «Коротко о решении»: слева техническая суть, "
                 "справа маркетинговая. Пример детекции — реальный кадр с рамками модели.")


def slide_concept(s):
    label = shape(s, "Прямоугольник 14")
    set_text_keep_style(label, "КОНЦЕПЦИЯ И АРХИТЕКТУРА")
    label.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xD6, 0xE4)
    fill(shape(s, "Текст 2").text_frame, [
        ("Портфель на карте. ", "Все площадки на одной карте: цвет маркера — проблемы, вопросы или норма."),
        ("План → снимок → сигнал. ", "Объект заводится с календарным планом, снимки распознаёт YOLO, "
                                     "техника сверяется с этапом на дату кадра."),
        ("Прогноз и динамика. ", "Оценка отставания по серии снимков, признаки простоя (R-08), "
                                 "зоны кадра: опасная зона и склад (R-09, R-10)."),
        ("Инспектор в контуре. ", "Стол разбора сигналов «верно / ошибка»; вердикты копятся "
                                  "в статистику качества распознавания."),
        ("ИИ-помощник. ", "Отвечает по данным объекта и предлагает перенос дат только с подтверждением."),
    ], size=11, space=5)
    pic = next(sh for sh in s.shapes if sh.shape_type == 13)
    replace_picture(s, pic, SHOTS / "arch_simple_crop.png")
    set_notes(s, "Концепция и архитектура. Правила: R-01…R-03 — техника и этап плана, "
                 "R-08 — вероятный простой по серии снимков, R-09 — техника в опасной зоне, "
                 "R-10 — кран в зоне склада на этапе монтажа.")


def slide_tech(s):
    fill(shape(s, "Текст 2").text_frame, [
        ("Клиент-сервер: ", "Next.js 15 — интерфейс, FastAPI — API, правила и прогноз, отдельный worker детекции."),
        ("Детектор: ", "YOLOv8m 1280 px с нарезкой SAHI; датасет из открытых наборов, размечен VLM и вручную."),
        ("ИИ-помощник: ", "GLM-5.3 Flash через OpenCode Zen."),
        ("Почему так: ", "общий Python-стек для CV и API, детекция работает и без облачного GPU."),
    ], size=10, space=3)
    pic = next(sh for sh in s.shapes if sh.shape_type == 13 and sh.width > Inches(10))
    replace_picture(s, pic, SHOTS / "s7_20_new.png")
    set_notes(s, "Обоснование технологий. YOLOv8m выбран как баланс точности и скорости на CPU/GPU; "
                 "SAHI нужен для мелкой техники на обзорных кадрах. Помощник — GLM-5.3 Flash.")


def slide_main_screen(s):
    picture_fit(s, SHOTS / "crop_home.png", 0.6, 2.05, 7.0, 3.8)
    tb = textbox(s, 7.75, 2.1, 4.7, 3.7)
    fill(tb.text_frame, [
        ("Карта Москвы: ", "на маркере число сигналов, цвет — проблемы, вопросы или норма."),
        ("Карточка объекта сбоку: ", "этап по плану и факт на снимке, прогноз отставания, техника этапа."),
        ("«Начать разбор»: ", "очередь сигналов, которые ждут решения инспектора."),
        ("«Спросить BuildWatch»: ", "ИИ-помощник по всему портфелю."),
        ("Фильтры и список: ", "по типу объекта и статусу."),
    ], size=11.5, space=6, heading="Главный экран — портфель объектов", heading_size=15)
    set_notes(s, "Скриншот рабочего прототипа, 29.09.2026. Демо-данные: 6 объектов в Москве.")


def slide_object_card(s):
    set_text_keep_style(shape(s, "Прямоугольник 1"), "Пример работы. Карточка объекта и отчёт")
    picture_fit(s, SHOTS / "crop_object.png", 0.65, 2.05, 4.9, 3.8)
    picture_fit(s, SHOTS / "crop_report.png", 5.6, 2.05, 5.2, 3.8)
    tb = textbox(s, 10.95, 2.1, 1.6, 3.7)
    steps = ["Снимок", "Детекция YOLO", "Этап на дату кадра", "Правило → сигнал", "Вердикт инспектора",
             "Прогноз и отчёт"]
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    add_run(tf.paragraphs[0], "От кадра до результата", 11, PURPLE, bold=True)
    tf.paragraphs[0].space_after = Pt(6)
    for i, step in enumerate(steps, 1):
        p = tf.add_paragraph()
        add_run(p, f"{i}  ", 10, ACCENT, bold=True)
        add_run(p, step, 10, DARK)
        p.space_after = Pt(5)
    cap = textbox(s, 0.6, 6.1, 12.0, 0.9)
    fill(cap.text_frame, [
        ("Карточка: ", "статус и прогноз, этапы плана, снимок с рамками YOLO и зонами, план/факт, "
                       "активность техники, сигналы. "),
        ("Отчёт: ", "прогноз, динамика за 30 дней, качество распознавания по вердиктам, план и снимки; печать в PDF."),
    ], size=10, color=WHITE, lead_color=WHITE, space=2)
    set_notes(s, "Скриншоты рабочего прототипа, объект «ЖК Северный, корп. 12». "
                 "Показан весь путь: распознавание техники → сопоставление с этапом → сигнал → отчёт.")


def slide_scaling(s):
    boxes = sorted((sh for sh in s.shapes if sh.name.startswith("Прямоугольник: скругленные углы")),
                   key=lambda sh: sh.left)
    horizontal = [
        ("База знаний по объекту: ", "проектная и исполнительная документация, акты, предписания."),
        ("Подключение планов: ", "импорт графиков из MS Project, Primavera и систем ДГП."),
        ("Сообщения и переписка: ", "сигнал связывается с чатом подрядчика, уведомления в Telegram."),
        ("Сметы и объёмы: ", "сверка техники и видов работ со сметой и ресурсной ведомостью."),
        ("Масштаб портфеля: ", "PostgreSQL, пул воркеров, роли для ДГП, заказчиков и подрядчиков."),
    ]
    vertical = [
        ("Качество распознавания: ", "дообучение на московских кадрах по вердиктам инспекторов."),
        ("Реальное время: ", "RTSP-потоки, трекинг техники, учёт моточасов и простоев."),
        ("Новые классы: ", "башенный кран, сваебой, асфальтоукладчик, люди и СИЗ."),
        ("Геометрия площадки: ", "маска обзора камеры, калибровка зон по стройгенплану."),
        ("Оценка прогресса: ", "несколько камер на объект, объёмы работ и этажность по кадрам."),
    ]
    for box, (title, subtitle, items) in zip(boxes, [
        ("Горизонтальное развитие", "больше функций вокруг объекта", horizontal),
        ("Вертикальное развитие", "глубже и точнее распознавание", vertical),
    ]):
        x = box.left / 914400 + 0.3
        tb = textbox(s, x, box.top / 914400 + 0.25, box.width / 914400 - 0.6, box.height / 914400 - 0.4)
        tf = tb.text_frame
        fill(tf, items, size=11.5, space=6, heading=title, heading_size=17, subheading=subtitle)
    arrow_h = textbox(s, 0.8, 6.15, 5.2, 0.35)
    add_run(arrow_h.text_frame.paragraphs[0], "→  шире: новые данные и сценарии", 10, WHITE, bold=True)
    arrow_v = textbox(s, 7.1, 6.15, 5.2, 0.35)
    add_run(arrow_v.text_frame.paragraphs[0], "↑  глубже: точность и скорость CV", 10, WHITE, bold=True)
    set_notes(s, "Две оси развития. Горизонталь — больше функций и данных вокруг объекта. "
                 "Вертикаль — качество и скорость распознавания.")


def main():
    prs = Presentation(str(DECK))
    slides = list(prs.slides)
    slide_team(slides[2])
    slide_solution(slides[4])
    slide_concept(slides[5])
    slide_tech(slides[6])
    slide_main_screen(slides[7])
    slide_object_card(slides[8])
    slide_scaling(slides[10])
    set_notes(slides[9], "Схема обработки кадра: SAHI-нарезка, YOLO, сборка рамок, сопоставление с этапом плана "
                         "на дату снимка и отображение в карточке.")
    set_notes(slides[11], "Рекомендации по установке камер: две камеры — обзорная A и рабочая B; "
                          "требования к видимости техники в кадре.")
    for i, s in enumerate(slides[5:], 6):
        for sh in s.shapes:
            if sh.name == "Прямоугольник 13" and sh.has_text_frame:
                set_text_keep_style(sh, f"{i:02d}")
    prs.save(str(DECK))
    print("saved", DECK)


if __name__ == "__main__":
    main()
