"""Preview dated CSV/XLSX construction plans before they replace an object plan."""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from zipfile import BadZipFile

from fastapi import HTTPException

HEADER_ALIASES = {
    "name": ("этап", "вид работ", "наименование", "работа", "stage", "work"),
    "date_from": ("дата начала", "начало", "начало работ", "date from", "start"),
    "date_to": ("дата окончания", "окончание", "конец", "завершение", "date to", "finish", "end"),
}


def _header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _date(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    raw = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, pattern).date().isoformat()
        except ValueError:
            pass
    raise ValueError("Дата должна быть в формате ДД.ММ.ГГГГ или ГГГГ-ММ-ДД")


def _kind(name: str) -> str:
    text = name.lower()
    if any(word in text for word in ("котлован", "землян", "фундамент", "сва", "шпунт")):
        return "excavation"
    if any(word in text for word in ("каркас", "монолит", "стен", "перекрыт")):
        return "frame"
    if any(word in text for word in ("фасад", "окн", "двер")):
        return "facade"
    if "кровл" in text or "крыши" in text:
        return "roof"
    if any(word in text for word in ("подготов", "территор", "огражден")):
        return "ground"
    return "other"


def parse_plan(content: bytes, filename: str) -> dict:
    if not content or len(content) > 2 * 1024 * 1024:
        raise HTTPException(400, "Выберите CSV или XLSX размером до 2 МБ")
    filename = filename.lower()
    try:
        if filename.endswith(".csv"):
            try:
                decoded = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                decoded = content.decode("cp1251")
            sample = decoded[:4096]
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
            except csv.Error:
                delimiter = ";"
            rows = list(csv.reader(io.StringIO(decoded), delimiter=delimiter))
        elif filename.endswith(".xlsx"):
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True, max_row=510))
            finally:
                workbook.close()
        else:
            raise HTTPException(400, "Поддерживаются файлы CSV и XLSX")
    except (UnicodeError, csv.Error, ValueError, OSError, BadZipFile) as exc:
        raise HTTPException(400, "Не удалось прочитать таблицу") from exc

    header_row = None
    columns = {}
    for index, row in enumerate(rows[:15]):
        found = {
            field: next((col for col, value in enumerate(row) if _header(value) in aliases), None)
            for field, aliases in HEADER_ALIASES.items()
        }
        if all(value is not None for value in found.values()):
            header_row, columns = index, found
            break
    if header_row is None:
        raise HTTPException(422, "Нужны столбцы «Этап», «Дата начала», «Дата окончания». Справочник без дат не является календарным планом.")

    stages = []
    issues = []
    for line, row in enumerate(rows[header_row + 1:510], start=header_row + 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        try:
            name = str(row[columns["name"]] or "").strip()
            if not name:
                raise ValueError("нет названия этапа")
            start = _date(row[columns["date_from"]])
            end = _date(row[columns["date_to"]])
            if end < start:
                raise ValueError("окончание раньше начала")
            stages.append({"name": name[:200], "date_from": start, "date_to": end,
                           "kind": _kind(name), "status": "future"})
        except (IndexError, ValueError) as exc:
            issues.append(f"Строка {line}: {exc}")
    if not stages:
        raise HTTPException(422, "В таблице нет этапов с корректными датами")
    return {"stages": stages, "issues": issues[:20], "total_rows": len(stages) + len(issues)}
