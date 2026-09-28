"""Import the user-provided DGP work catalog and photos into the public demo DB."""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from PIL import Image
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from backend import db  # noqa: E402
from backend.detector_config import DEFAULT_MODEL, detector_available  # noqa: E402
from backend.evaluate import evaluate_object  # noqa: E402

OBJECT_NAME = 'Демо-набор ДГП: строительные снимки'
LEGACY_NAME = 'Синтетическая демонстрационная площадка'
SAMPLES = ROOT / 'context/dgp_extract/samples'
CATALOG = ROOT / 'context/dgp_extract/Сводный перечень строительных работ_ЛТЦ.xlsx'
OBJECT_DESCRIPTION = (
    'Материалы из пользовательского архива 7.ДГП_датасеты.zip: 100 строительных снимков '
    'и справочник видов работ. Детектор equipment распознаёт технику на снимках; '
    'календарный план демо-этапов задан для правил «этап → техника». '
    'Точные даты съёмки и геопривязка источником не предоставлены.'
)

# Demo calendar: current date falls into excavation so R-03 (dump truck) and
# crane-vs-stage rules have something to evaluate against live detections.
DEMO_STAGES: list[tuple[str, str, str, str]] = [
    ('ground', 'Обустройство строительной площадки', '2026-07-01', '2026-08-15'),
    ('excavation', 'Разработка грунта с креплением котлована', '2026-08-16', '2026-10-15'),
    ('frame', 'Монолитные работы ниже отм. 0', '2026-10-16', '2026-12-31'),
]


def import_catalog(con) -> int:
    sheet = load_workbook(CATALOG, data_only=True, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[2]
    type_columns = [
        (index, value.strip()) for index, value in enumerate(headers)
        if isinstance(value, str) and value.strip() not in {'№ п/п', 'Вид работ'}
    ]
    con.execute('DELETE FROM catalog')
    count = 0
    for row in rows[3:]:
        code, name = row[0], row[1]
        if not isinstance(name, str) or not name.strip():
            continue
        applies = [
            category for index, category in type_columns
            if str(row[index]).strip() in {'˅', 'v', 'V', '✓'}
        ]
        con.execute(
            'INSERT INTO catalog(code,name,applies) VALUES (?,?,?)',
            (str(code or '').strip(), name.strip(), ','.join(applies)),
        )
        count += 1
    return count


def main() -> None:
    if db.DATABASE_URL:
        raise RuntimeError('Public DGP demo seeding is supported for SQLite only')
    if not CATALOG.is_file():
        raise FileNotFoundError(f'DGP work catalog is missing: {CATALOG}')
    images = sorted(SAMPLES.glob('Screenshot_*.png'), key=lambda path: int(path.stem.split('_')[-1]))
    if len(images) != 100:
        raise RuntimeError(f'Expected 100 supplied DGP images, found {len(images)}')

    db.init()
    con = db.connect()
    try:
        row = con.execute('SELECT id FROM objects WHERE name=?', (OBJECT_NAME,)).fetchone()
        if row is None:
            row = con.execute('SELECT id FROM objects WHERE name=?', (LEGACY_NAME,)).fetchone()
        if row is None:
            cursor = con.execute(
                'INSERT INTO objects(name,type,district,address,description) VALUES (?,?,?,?,?)',
                (OBJECT_NAME, 'Жильё', 'Демо', 'Точный адрес и координаты в наборе не указаны',
                 OBJECT_DESCRIPTION),
            )
            object_id = cursor.lastrowid
        else:
            object_id = row['id']
            con.execute(
                'UPDATE objects SET name=?,type=?,district=?,address=?,description=? WHERE id=?',
                (OBJECT_NAME, 'Жильё', 'Демо', 'Точный адрес и координаты в наборе не указаны',
                 OBJECT_DESCRIPTION, object_id),
            )

        snapshot_ids = [r['id'] for r in con.execute('SELECT id FROM snapshots WHERE object_id=?', (object_id,))]
        if snapshot_ids:
            con.execute('DELETE FROM warnings WHERE object_id=?', (object_id,))
            con.execute('DELETE FROM detection_jobs WHERE snapshot_id IN (SELECT id FROM snapshots WHERE object_id=?)', (object_id,))
            con.execute('DELETE FROM detections WHERE snapshot_id IN (SELECT id FROM snapshots WHERE object_id=?)', (object_id,))
            con.execute('DELETE FROM snapshots WHERE object_id=?', (object_id,))
        con.execute('DELETE FROM stages WHERE object_id=?', (object_id,))
        catalog_count = import_catalog(con)

        for position, (kind, name, date_from, date_to) in enumerate(DEMO_STAGES):
            con.execute(
                'INSERT INTO stages(object_id, position, kind, name, date_from, date_to, status) '
                'VALUES (?,?,?,?,?,?,?)',
                (
                    object_id,
                    position,
                    kind,
                    name,
                    date_from,
                    date_to,
                    'current' if kind == 'excavation' else ('done' if kind == 'ground' else 'future'),
                ),
            )

        imported_on = date.today().isoformat()
        for image_path in images:
            with Image.open(image_path) as image:
                width, height = image.size
            con.execute(
                'INSERT INTO snapshots(object_id,filename,src,captured_at,status,width,height) '
                'VALUES (?,?,?,?,?,?,?)',
                (object_id, image_path.name, 'samples', imported_on, 'new', width, height),
            )
        con.commit()
    finally:
        con.close()

    # Queue live detection after commit so the worker sees durable rows.
    queued = 0
    if detector_available(DEFAULT_MODEL):
        con = db.connect()
        try:
            for row in con.execute(
                'SELECT id FROM snapshots WHERE object_id=? ORDER BY id', (object_id,)
            ):
                snap_id = int(row['id'])
                con.execute('UPDATE snapshots SET status=? WHERE id=?', ('processing', snap_id))
                con.execute(
                    "INSERT INTO detection_jobs(snapshot_id,model,status,attempts,error,started_at,finished_at)"
                    " VALUES (?,?,'queued',0,NULL,NULL,NULL)"
                    " ON CONFLICT(snapshot_id) DO UPDATE SET model=excluded.model,status='queued',"
                    " attempts=0,error=NULL,started_at=NULL,finished_at=NULL",
                    (snap_id, DEFAULT_MODEL),
                )
                queued += 1
            con.commit()
        finally:
            con.close()

    evaluate_object(object_id)
    print(
        f'Imported object {object_id}: {len(images)} images, {catalog_count} work catalog entries, '
        f'{len(DEMO_STAGES)} stages, queued={queued}'
    )


if __name__ == '__main__':
    main()
