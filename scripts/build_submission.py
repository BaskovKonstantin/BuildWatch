"""Build the LCT submission deck and PDF from the supplied template and docs."""
from __future__ import annotations

import argparse
import re
from pathlib import Path
from xml.sax.saxutils import escape

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = Path('/mnt/d/Downloads/ЛЦТ2026 Шаблон презентации.pptx')
DECK = ROOT / 'docs/BuildWatch-ЛЦТ2026.pptx'
MARKDOWN = ROOT / 'docs/submission.md'
PDF = ROOT / 'docs/submission.pdf'

PURPLE = RGBColor(49, 15, 83)
PINK = RGBColor(255, 0, 83)
LILAC = RGBColor(138, 131, 209)
PALE = RGBColor(255, 214, 228)
INK = RGBColor(28, 29, 34)
WHITE = RGBColor(255, 255, 255)


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.03)
    frame.margin_top = frame.margin_bottom = 0
    para = frame.paragraphs[0]
    if align is not None:
        para.alignment = align
    run = para.add_run()
    run.text = value
    run.font.name = 'Montserrat'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add(slides, title, label, number, dark=False):
    slide = slides.add_slide(slides._parent.slide_layouts[22])
    rect(slide, 0, 0, 13.333, 7.5, PURPLE if dark else WHITE)
    text(slide, label.upper(), .55, .34, 9, .28, 10, PALE if dark else LILAC, True)
    text(slide, title, .55, .82, 12.2, .95, 32, WHITE if dark else PURPLE, True)
    text(slide, str(number).zfill(2), 12.2, 7.02, .55, .23, 10, PALE if dark else LILAC)
    return slide


def card(slide, title, body, x, y, w, h, dark=False):
    rect(slide, x, y, w, h, PURPLE if dark else RGBColor(249, 247, 252), True)
    text(slide, title, x+.25, y+.24, w-.5, .48, 18, PALE if dark else PURPLE, True)
    text(slide, body, x+.25, y+.82, w-.5, h-1, 15, WHITE if dark else INK)


def make_deck(template: Path):
    prs = Presentation(template)
    # Add before deleting guidance slides so new package part names cannot collide.
    cover = prs.slides[6]
    rect(cover, 0, 0, 13.333, 7.5, PURPLE)
    rect(cover, 8.85, 0, 4.48, 7.5, PINK)
    text(cover, 'BUILDWATCH', .7, 1.0, 8.0, .8, 38, WHITE, True)
    text(cover, 'Снимок. План. Сигнал.', .7, 2.15, 8.1, .6, 26, PALE, True)
    text(cover, 'Автоматизированный контроль строительных площадок Москвы', .7, 3.1, 7.7, 1.4, 28, WHITE)
    text(cover, 'ЛЦТ 2026 · задача ДГП Москвы', .7, 6.45, 7, .35, 15, WHITE)
    text(cover, 'ПРОТОТИП', 9.15, 3.0, 3.9, .8, 27, WHITE, True, PP_ALIGN.CENTER)

    s = add(prs.slides, 'Задача и ответ', 'Коротко о решении', 2)
    card(s, 'Проблема', 'Ручной просмотр камер и сверка с календарным планом занимают время и не дают своевременных сигналов.', .55, 2.05, 5.9, 3.9)
    card(s, 'Решение', 'BuildWatch связывает датированный снимок, обнаруженную технику и этап работ. Правила формируют объяснимый сигнал инспектору.', 6.8, 2.05, 5.95, 3.9, True)
    text(s, 'Один кадр — повод для проверки, а не доказательство простоя.', .72, 6.48, 11.8, .45, 18, PINK, True)

    s = add(prs.slides, 'Команда и контакты', 'Обязательный блок шаблона', 3)
    card(s, 'Команда BuildWatch', 'Капитан: [ФИО, специальность]\nУчастники и роли: [заполнить]\nГород, организация: [заполнить]', .55, 2.0, 6.0, 4.5)
    card(s, 'Контакт для жюри', 'Имя: [заполнить]\nПочта: [заполнить]\nТелефон / мессенджер: [заполнить]', 6.8, 2.0, 5.95, 4.5, True)

    s = add(prs.slides, 'Суть и уникальность', 'Обязательный блок шаблона', 4)
    card(s, 'Проверка по плану', 'Справочник работ ЛТЦ задаёт виды этапов; даты берутся из отдельного календарного плана объекта.', .55, 2.05, 3.86, 3.85)
    card(s, 'Проверяемый вывод', 'У каждого сигнала есть правило, объяснение, снимок и возможность подтвердить или отклонить вывод.', 4.73, 2.05, 3.86, 3.85, True)
    card(s, 'Честная неопределённость', 'Ошибка детектора, закрытая зона камеры или неполный график не маскируются под доказанное нарушение.', 8.9, 2.05, 3.86, 3.85)

    s = add(prs.slides, 'Почему эта задача и что было сложно', 'Команда и вызовы', 5)
    card(s, 'Мотивация', 'Камеры уже дают регулярные наблюдения. Связь снимка с планом может ускорить проверку объектов инспектором.', .55, 2.0, 5.95, 4.15)
    card(s, 'Вызовы разработки', 'Разные ракурсы и мелкая техника; неоднозначные классы; справочник ЛТЦ без календарных дат; отсутствие официальной разметки для проверки модели.', 6.8, 2.0, 5.95, 4.15, True)
    text(s, 'Команда разделила план, детекцию и правила, а спорные выводы передала на проверку человеку.', .65, 6.55, 11.95, .42, 16, PINK, True)

    s = add(prs.slides, 'Что уже работает', 'Прототип', 6)
    items = [('01', 'Объекты и план'), ('02', 'Загрузка снимка'), ('03', 'Очередь и детекция'), ('04', 'Правила отклонений'), ('05', 'Проверка инспектором'), ('06', 'Отчёт объекта')]
    for i, (num, value) in enumerate(items):
        x = .55 + (i % 3) * 4.17
        y = 2.05 + (i // 3) * 2.05
        rect(s, x, y, 3.88, 1.72, PURPLE if i in (2, 3) else RGBColor(249, 247, 252), True)
        color = WHITE if i in (2, 3) else PURPLE
        text(s, num, x+.2, y+.17, .62, .44, 21, PINK, True)
        text(s, value, x+.2, y+.78, 3.43, .64, 19, color, True)

    s = add(prs.slides, 'От камеры до сигнала', 'Связь снимка и графика', 7)
    flow = [('Снимок', 'ID объекта\nдата и время'), ('Детекция', 'класс · рамка\nуверенность'), ('План', 'этап на дату\nожидаемые классы'), ('Правило', 'R-01 / R-02\nR-03 / R-07')]
    for i, (heading, body) in enumerate(flow):
        card(s, heading, body, .55+i*3.16, 2.32, 2.83, 3.16, i == 3)
        if i < 3:
            text(s, '→', 3.39+i*3.16, 3.52, .28, .55, 22, PINK, True)
    text(s, 'Если детекция ещё выполняется или завершилась ошибкой, отсутствие техники не выводится.', .65, 6.13, 12.1, .55, 17, PURPLE)

    s = add(prs.slides, 'Архитектура', 'Техническая реализация', 8)
    card(s, 'Next.js', 'Портфель, карточка, план, снимки, рамки и отчёт.', .55, 2.1, 3.8, 3.7)
    card(s, 'FastAPI + БД', 'API, валидация, SQLite для демо и схема PostgreSQL.', 4.77, 2.1, 3.8, 3.7, True)
    card(s, 'Worker + YOLO', 'Очередь, локальный инференс, детекции и пересчёт правил.', 8.99, 2.1, 3.8, 3.7)
    text(s, 'Разделение API и инференса сохраняет отзывчивость интерфейса.', .66, 6.27, 11.9, .47, 18, PINK, True)

    s = add(prs.slides, 'Методика предупреждений', 'Правила и объяснимость', 9)
    for i, (head, body) in enumerate([
        ('R-01', 'Лишняя для этапа техника'), ('R-02', 'Низкая уверенность / конфликт'),
        ('R-03', 'Нет обязательной техники'), ('R-07', 'Кран на фасадном этапе')]):
        card(s, head, body, .55+(i%2)*6.25, 2.0+(i//2)*2.18, 5.95, 1.88, i == 2)
    text(s, 'Для котлована ожидаются экскаватор и самосвал. Сигнал остаётся предметом проверки.', .65, 6.6, 11.9, .42, 16, PURPLE)

    s = add(prs.slides, 'Прототип в работе', 'Демонстрация', 10)
    # Recreate the interface as editable shapes; no organizer-owned photo is embedded.
    rect(s, .55, 1.85, 8.1, 4.95, RGBColor(247, 246, 250), True)
    text(s, 'ЖК «Северный» · демонстрационный объект', .82, 2.1, 7.3, .48, 17, PURPLE, True)
    rect(s, .8, 2.75, 2.18, 3.55, WHITE, True)
    text(s, 'ПЛАН РАБОТ', .98, 2.97, 1.7, .33, 11, LILAC, True)
    text(s, '1  Подготовка\n\n2  Котлован\n\n3  Каркас\n\n4  Фасад', .98, 3.49, 1.83, 2.36, 14, INK)
    rect(s, 3.2, 2.75, 3.22, 3.55, PALE, True)
    text(s, 'СИНТЕТИЧЕСКИЙ КАДР', 3.42, 2.97, 2.8, .35, 11, PURPLE, True)
    rect(s, 3.65, 4.12, 1.2, .86, LILAC, True)
    rect(s, 4.75, 4.51, .86, .46, PURPLE, True)
    text(s, 'экскаватор · 0,83', 3.43, 5.45, 2.72, .35, 12, PURPLE, True)
    rect(s, 6.62, 2.75, 1.75, 3.55, WHITE, True)
    text(s, 'СИГНАЛ', 6.8, 2.97, 1.36, .3, 11, LILAC, True)
    text(s, 'R-03\n\nНе виден\nсамосвал\n\nПроверить', 6.8, 3.43, 1.35, 2.45, 13, PURPLE, True)
    card(s, 'Карточка объекта', 'Слева — этапы плана. В центре — изображение и рамки модели. Справа — сигналы. На схеме использованы синтетические данные.', 8.95, 2.05, 3.84, 4.32, True)

    s = add(prs.slides, 'Данные и ограничения модели', 'Оценка качества', 11)
    card(s, 'Данные', '8 целевых классов ТЗ. Снимки ДГП не являются размеченным официальным тестом. Внешние наборы имеют разные лицензии и домены.', .55, 2.0, 5.95, 4.2)
    card(s, 'Оценка', 'Ранняя equipment_v2: mAP@50 = 0,786 на внешнем holdout. Для текущих весов v6 нужна отдельная замороженная оценка по всем классам на московских камерах.', 6.8, 2.0, 5.95, 4.2, True)
    text(s, 'Метрики другой модели или другого набора не переносятся на текущий прототип.', .65, 6.55, 11.95, .45, 17, PINK, True)

    s = add(prs.slides, 'Камеры: условия достоверности', 'Рекомендации для площадки', 12)
    card(s, 'Размещение', 'Несколько фиксированных ракурсов: рабочие зоны и маршруты техники, минимум перекрытий.', .55, 2.0, 5.95, 1.95)
    card(s, 'Качество кадра', 'Достаточное разрешение для удалённых объектов, свет ночью, контроль осадков и чистоты объектива.', 6.8, 2.0, 5.95, 1.95, True)
    card(s, 'Метаданные', 'Единое время съёмки и идентификатор площадки, синхронизация с календарным графиком.', .55, 4.24, 5.95, 1.95, True)
    card(s, 'Проверка отсутствия', 'Прежде чем выдавать сигнал, убедиться, что нужная техника могла попасть в поле зрения.', 6.8, 4.24, 5.95, 1.95)

    s = add(prs.slides, 'Следующие шаги', 'Развитие решения', 13, True)
    for i, (head, body) in enumerate([
        ('01 · Качество', 'Размеченный московский тест, метрики по каждому классу и погоде.'),
        ('02 · Мелкие объекты', 'Тайлинг SAHI и оценка задержки на слабом сервере.'),
        ('03 · Аналитика', 'Несколько камер, учёт видимости и временные ряды.'),
        ('04 · Внедрение', 'Пилот с инспекторами и калибровка правил по площадкам.')]):
        card(s, head, body, .55+(i%2)*6.25, 1.95+(i//2)*2.35, 5.95, 2.08, True)

    s = add(prs.slides, 'BuildWatch', 'Итог', 14, True)
    text(s, 'Снимок → этап → объяснимый сигнал', .75, 2.14, 11.85, .86, 30, WHITE, True)
    text(s, 'Работающий интерфейс, правила и проверяемый человеком результат.', .78, 3.35, 11.8, .8, 22, PALE)
    text(s, 'Репозиторий и инструкция по запуску — в README.', .78, 5.12, 11.7, .53, 18, WHITE)
    ids = prs.slides._sldIdLst
    title_id = ids[6]
    for sid in list(ids)[:37]:
        if sid is title_id:
            continue
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    prs.core_properties.title = 'BuildWatch — ЛЦТ 2026'
    prs.core_properties.subject = 'Автоматизированный контроль строительных площадок'
    prs.save(DECK)


def make_pdf():
    font_pairs = [
        (Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')),
        (Path('C:/Windows/Fonts/arial.ttf'), Path('C:/Windows/Fonts/arialbd.ttf')),
        (Path('/mnt/c/Windows/Fonts/arial.ttf'), Path('/mnt/c/Windows/Fonts/arialbd.ttf')),
    ]
    font, bold = next(((regular, heavy) for regular, heavy in font_pairs if regular.exists() and heavy.exists()), (None, None))
    if font is None:
        raise FileNotFoundError('A Cyrillic DejaVu Sans or Arial font is required to build the PDF')
    pdfmetrics.registerFont(TTFont('DejaVu', str(font)))
    pdfmetrics.registerFont(TTFont('DejaVuBold', str(bold)))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='CoverBW', fontName='DejaVuBold', fontSize=19, leading=26, textColor=colors.HexColor('#310F53'), spaceAfter=13))
    styles.add(ParagraphStyle(name='HeadingBW', fontName='DejaVuBold', fontSize=12, leading=18, textColor=colors.HexColor('#310F53'), spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle(name='BodyBW', fontName='DejaVu', fontSize=9, leading=14, spaceAfter=6))
    styles.add(ParagraphStyle(name='SmallBW', fontName='DejaVu', fontSize=8, leading=12, leftIndent=11, spaceAfter=3))
    story = []
    for raw in MARKDOWN.read_text(encoding='utf-8').splitlines():
        line = raw.strip()
        if not line or line.startswith('```') or line.startswith('|---'):
            if not line:
                story.append(Spacer(1, 2*mm))
            continue
        if line.startswith('# '):
            style = styles['CoverBW']; line = line[2:]
        elif line.startswith('## '):
            style = styles['HeadingBW']; line = line[3:]
        elif line.startswith('|'):
            style = styles['SmallBW']; line = line.strip('|').replace('|', '  ·  ')
        elif line.startswith(('- ', '1. ', '2. ', '3. ', '4. ', '5. ')):
            style = styles['SmallBW']
        else:
            style = styles['BodyBW']
        line = escape(line)
        line = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', line)
        line = re.sub(r'`([^`]+)`', r'<font color="#520978">\1</font>', line)
        line = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', line)
        story.append(Paragraph(line, style))
    doc = SimpleDocTemplate(str(PDF), pagesize=(210*mm, 297*mm), leftMargin=20*mm, rightMargin=20*mm, topMargin=19*mm, bottomMargin=18*mm,
                            title='BuildWatch — документация ЛЦТ 2026')
    doc.build(story)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--template', type=Path, default=DEFAULT_TEMPLATE,
                        help='Path to the supplied LCT2026 PPTX template')
    args = parser.parse_args()
    make_deck(args.template)
    make_pdf()
    print(DECK)
    print(PDF)
