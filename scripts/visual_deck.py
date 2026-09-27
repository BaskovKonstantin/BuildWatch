"""Build the visual LCT deck while preserving mandatory template slides 7–11."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs/BuildWatch-ЛЦТ2026.pptx"
ASSETS = ROOT / "docs/assets"
PURPLE = RGBColor(49, 15, 83)
PINK = RGBColor(255, 0, 83)
LILAC = RGBColor(138, 131, 209)
INK = RGBColor(28, 29, 34)
WHITE = RGBColor(255, 255, 255)
PALE = RGBColor(249, 247, 252)


def fill(shape, *lines: str) -> None:
    """Replace text in existing paragraphs, retaining each template run's style."""
    for index, paragraph in enumerate(shape.text_frame.paragraphs):
        value = lines[index] if index < len(lines) else ""
        if paragraph.runs:
            paragraph.runs[0].text = value
            for run in paragraph.runs[1:]:
                run.text = ""
        elif value:
            paragraph.add_run().text = value


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = 0
    run = frame.paragraphs[0].add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def shape(slide, x, y, w, h, color, rounded=False, outline=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    box = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = color
    if outline:
        box.line.color.rgb = outline
    else:
        box.line.fill.background()
    return box


def visual_slide(prs, eyebrow, title, page, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[22])
    shape(slide, 0, 0, 13.333, 7.5, PURPLE if dark else WHITE)
    text(slide, eyebrow.upper(), .55, .33, 11.7, .32, 10, PINK if dark else LILAC, True)
    text(slide, title, .55, .8, 12, .8, 31, WHITE if dark else PURPLE, True)
    text(slide, f"{page:02}", 12.35, 7.02, .45, .25, 10, WHITE if dark else LILAC)
    return slide


def image(slide, path: Path, x, y, w):
    if not path.exists():
        raise FileNotFoundError(f"Presentation screenshot is missing: {path}")
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))


def mandatory_slides(prs):
    # Only text values change; all original shapes, geometry, fills and art stay.
    s = prs.slides[6]
    fill(s.shapes[0], "BUILDWATCH")
    fill(s.shapes[1], "Снимок → план → сигнал")
    fill(s.shapes[2], "Сервис контроля стройплощадок Москвы · ЛЦТ 2026")

    s = prs.slides[7]
    fill(s.shapes[6], "BUILDWATCH")
    fill(s.shapes[11], "Снимок и этап работ превращаются в проверяемый сигнал инспектору.")
    fill(s.shapes[3], "План, детекция и правила дают объяснение каждого сигнала.")
    fill(s.shapes[4], "Капитан: [ФИО, специальность]", "Кол-во участников: [заполнить]",
         "Краткое описание: [заполнить]", "История команды: [заполнить]",
         "Место работы/учёбы: [заполнить]", "Город и регион: [заполнить]")

    # The team slide has five cards in the original. Keep all for owner input.
    s = prs.slides[8]
    fill(s.shapes[20], "СОСТАВ КОМАНДЫ · ЗАПОЛНИТЬ ПЕРЕД СДАЧЕЙ")

    s = prs.slides[9]
    fill(s.shapes[2], "О КОМАНДЕ И ВЫЗОВАХ")
    fill(s.shapes[3], "[Как собралась команда и какие проекты делала вместе]")
    fill(s.shapes[7], "Камеры дают регулярные наблюдения, которые можно сверять с планом.")
    fill(s.shapes[5], "Мелкая техника и разные ракурсы осложняют детекцию.",
         "Для оценки модели нужен размеченный московский набор.")

    s = prs.slides[10]
    fill(s.shapes[4], "Next.js → FastAPI → очередь → YOLO. Правила R-01/R-02/R-03/R-07 сверяют технику с этапом. Инспектор подтверждает или отклоняет сигнал.")
    fill(s.shapes[5], "Пилот на объектах: несколько камер, калибровка правил и размеченный тест. Затем — интеграция с календарными графиками и отчётностью.", "")


def architecture(prs):
    s = visual_slide(prs, "архитектура ПО", "От снимка до решения инспектора", 6)
    stages = [
        ("01", "КАМЕРА", "Снимок + время", .65),
        ("02", "API", "FastAPI + БД", 3.8),
        ("03", "CV WORKER", "YOLO + очередь", 6.95),
        ("04", "ИНСПЕКТОР", "Next.js + отчёт", 10.1),
    ]
    for number, title, note, x in stages:
        shape(s, x, 2.25, 2.55, 2.18, PALE, True)
        text(s, number, x+.19, 2.48, .63, .53, 26, PINK, True)
        text(s, title, x+.19, 3.17, 2.2, .4, 18, PURPLE, True)
        text(s, note, x+.19, 3.69, 2.2, .38, 15, INK)
        if x < 10:
            text(s, "→", x+2.65, 2.99, .37, .46, 27, PINK, True)
    shape(s, 3.38, 5.04, 6.48, .77, PURPLE, True)
    text(s, "Этап плана + правила R-01 / R-02 / R-03 / R-07", 3.64, 5.27, 6.0, .31, 15, WHITE, True)
    text(s, "Сигнал содержит причину и остаётся на проверку человеком", .68, 6.49, 11.6, .35, 18, PURPLE, True)


def product(prs):
    s = visual_slide(prs, "скриншоты синтетического демо", "Прототип: портфель и карточка объекта", 7)
    image(s, ASSETS / "synthetic-home.png", .55, 1.84, 5.92)
    image(s, ASSETS / "synthetic-object.png", 6.85, 1.84, 5.92)
    text(s, "Объекты · сигналы · снимки", .69, 6.23, 5.7, .38, 18, PURPLE, True)
    text(s, "Этап · ожидаемая техника · вывод", 6.99, 6.23, 5.75, .38, 18, PURPLE, True)
    text(s, "Снимки экрана сделаны на синтетических данных; кадры ДГП не показаны.", .69, 6.88, 11.7, .28, 11, LILAC)


def model(prs):
    s = visual_slide(prs, "обучение CV модели", "Результат и границы оценки", 8)
    text(s, "8", .73, 2.15, 2.0, 1.35, 66, PURPLE, True)
    text(s, "целевых классов\nтехники по ТЗ", .75, 3.55, 3.0, .75, 19, INK)
    text(s, "0,836", 4.58, 2.15, 3.7, 1.35, 58, PINK, True)
    text(s, "mAP@50 v11 после 3 эпох\nна внешнем holdout", 4.62, 3.55, 3.7, .75, 19, INK)
    shape(s, 8.93, 2.08, 3.77, 2.57, PURPLE, True)
    text(s, "v6", 9.25, 2.4, 2.1, .73, 39, WHITE, True)
    text(s, "локально подключена\nк прототипу", 9.27, 3.37, 3.06, .8, 19, WHITE)
    text(s, "v11 продолжает обучение отдельно; это ещё не оценка рабочего прототипа.", .74, 5.35, 11.85, .6, 20, PURPLE, True)
    text(s, "Ранняя v2: mAP@50 0,786. Для v6 нужен размеченный тест на московских камерах.", .74, 6.1, 11.85, .45, 16, INK)


def assistant(prs):
    s = visual_slide(prs, "управление через ИИ", "«Спросить BuildWatch» помогает инспектору", 9, True)
    shape(s, .72, 2.02, 5.38, 3.96, WHITE, True)
    text(s, "Вопрос на естественном языке", 1.03, 2.36, 4.75, .51, 21, PURPLE, True)
    shape(s, 1.03, 3.19, 4.75, 1.05, PALE, True)
    text(s, "Что происходит на объекте и\nпочему выдан сигнал?", 1.23, 3.43, 4.35, .63, 18, INK)
    text(s, "→", 6.4, 3.35, .55, .62, 34, PINK, True)
    shape(s, 7.02, 2.02, 5.58, 3.96, WHITE, True)
    text(s, "Ответ с опорой на данные", 7.35, 2.36, 4.95, .51, 21, PURPLE, True)
    text(s, "План · снимки · детекции · правила", 7.35, 3.2, 4.83, .54, 17, INK)
    text(s, "Действия с данными требуют подтверждения пользователя.", 7.35, 4.13, 4.73, 1.05, 17, PURPLE, True)
    text(s, "ИИ объясняет и предлагает действия; решение о нарушении принимает инспектор.", .77, 6.57, 11.8, .44, 18, WHITE, True)


def final(prs):
    s = visual_slide(prs, "следующий шаг", "Пилот на московских камерах", 10)
    milestones = [("01", "Разметка", "Тест по 8 классам"),
                  ("02", "Калибровка", "Ракурсы и пороги"),
                  ("03", "Интеграция", "Планы и отчёты")]
    for i, (number, heading, note) in enumerate(milestones):
        x = .78 + i*4.17
        shape(s, x, 2.35, 3.6, 3.1, PURPLE if i == 1 else PALE, True)
        fg = WHITE if i == 1 else PURPLE
        text(s, number, x+.25, 2.68, 1.4, .83, 45, PINK, True)
        text(s, heading, x+.25, 3.82, 3.1, .52, 23, fg, True)
        text(s, note, x+.25, 4.55, 3.1, .43, 17, fg)
    text(s, "Один кадр даёт повод проверить. Несколько ракурсов и план дают контекст.", .8, 6.49, 11.65, .42, 18, PURPLE, True)


def make_deck(template: Path):
    prs = Presentation(template)
    mandatory_slides(prs)
    # Append visual explanation before dropping unused guidance slides.
    architecture(prs)
    product(prs)
    model(prs)
    assistant(prs)
    final(prs)
    ids = prs.slides._sldIdLst
    keep = {sid.rId for sid in list(ids)[6:11]}
    for sid in list(ids)[:37]:
        if sid.rId not in keep:
            prs.part.drop_rel(sid.rId)
            ids.remove(sid)
    prs.core_properties.title = "BuildWatch — ЛЦТ 2026"
    prs.core_properties.subject = "Автоматизированный контроль строительных площадок"
    prs.save(OUTPUT)
